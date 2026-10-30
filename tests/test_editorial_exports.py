"""P0 editorial template integration tests. Synthetic identities only."""
import io,time,uuid
from fastapi.testclient import TestClient
from app.main import app,product
from app.web_identity import token_hash
from app.services.application_packs.editorial import (
    validate_profile,cv_html,letter_html,from_user_cv_html,pdf_from_html,
    docx_from_html_data
)
from pypdf import PdfReader
from docx import Document

BASE="https://rolehint.example"
PROFILE={
  "name":"Jordan Example","headline":"Senior Product Manager | Growth & Operations",
  "city":"Barcelona, Spain","email":"jordan@example.com",
  "linkedin":"https://www.linkedin.com/in/jordanexample/",
  "skills":["Product discovery","B2C growth","Cross-functional leadership"],
  "experiences":[
    {"title":"Product Manager","company":"Example Learning","period":"2023 – 2026",
     "bullets":["Designed onboarding experiments for an education product with active learners.",
                "Worked with engineers and analysts to investigate retention and service quality."]},
    {"title":"Product Lead","company":"Sample Studio","period":"2021 – 2023",
     "bullets":["Built and validated operational workflows for internal support teams."]}
  ],
  "extras":[{"heading":"Method","value":"Research, experiment, evaluate, iterate"}],
  "approved":False
}
JOB={"job_id":"ashby:test","company":"Acme Labs","title":"Senior Product Manager",
     "location":"Spain remote","source_url":"https://jobs.ashbyhq.com/acme/job123"}

def session():
    u=product.ensure_user(telegram_user={"id":90000000+int(uuid.uuid4().int%10000000),
                                           "first_name":"Editorial Test"})
    token="editorial-"+uuid.uuid4().hex+uuid.uuid4().hex
    product.conn.execute("INSERT INTO web_sessions VALUES(?,?,?,?)",
                         (token_hash(token),u,time.time()+3600,time.time()))
    product.conn.commit()
    return u,token

def test_import_template_preserves_editorial_structure_and_no_script():
    profile=validate_profile(PROFILE)
    h=cv_html(profile,JOB,[])
    imported=from_user_cv_html(h)
    assert imported["name"]=="Jordan Example"
    assert imported["headline"]==PROFILE["headline"]
    assert len(imported["experiences"])==2
    assert imported["experiences"][0]["title"]=="Product Manager"
    assert imported["experiences"][0]["company"]=="Example Learning"
    assert imported["approved"] is False
    assert "two-column" not in h
    malicious=dict(PROFILE,name="<script>alert(1)</script>")
    safe=cv_html(malicious,JOB,[])
    assert "<script>alert" not in safe
    assert "&lt;script&gt;" in safe
    assert "grid-template-columns:120px 1fr 208px" in safe
    assert "border-radius:50%" in safe
    assert "text-align:justify" in safe

def test_one_page_pdf_and_docx_renderer():
    profile=validate_profile({**PROFILE,"approved":True})
    cv=cv_html(profile,JOB,[])
    pdf=pdf_from_html(cv)
    assert pdf.startswith(b"%PDF-")
    assert len(PdfReader(io.BytesIO(pdf)).pages)==1
    for kind in ("cv","letter"):
        raw=docx_from_html_data(profile,JOB,kind,[],
                                "I want to improve the onboarding journey.","")
        doc=Document(io.BytesIO(raw))
        all_text="\n".join([p.text for p in doc.paragraphs]+[p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs])
        assert "Jordan Example" in all_text
        if kind=="cv":assert any(len(table.columns)==3 for table in doc.tables)
    letter=letter_html(profile,JOB,[],"I value the team’s open experimentation.")
    assert len(PdfReader(io.BytesIO(pdf_from_html(letter))).pages)==1

def test_auth_approval_and_template_export_no_cross_user(monkeypatch):
    monkeypatch.setenv("CFA_WEB_APP_URL",BASE)
    u,token=session(); other,other_token=session()
    job=product.upsert_user_job(u,{**JOB,"source":"ashby","source_job_id":"test-"+uuid.uuid4().hex})
    jid=job["job_id"]
    path="/web-api/jobs/"+jid
    with TestClient(app,base_url=BASE) as client:
        assert client.get("/web-api/editorial-profile").status_code==401
        client.cookies.set("rh_session",token)
        p=client.put("/web-api/editorial-profile",headers={"Origin":BASE},
                     json=PROFILE)
        assert p.status_code==200,p.text
        assert p.json()["profile"]["approved"] is False
        assert client.post(path+"/documents/export",headers={"Origin":BASE},
            json={"kind":"cv","format":"pdf","approved":True}).status_code==409
        revised=dict(PROFILE,approved=True)
        assert client.put("/web-api/editorial-profile",headers={"Origin":BASE},
            json=revised).status_code==200
        # No letter export without explicit text and approval for this job.
        assert client.post(path+"/documents/export",headers={"Origin":BASE},
            json={"kind":"letter","format":"html","approved":True}).status_code==409
        assert client.put(path+"/documents/draft",headers={"Origin":BASE},
            json={"interest":"I care about the way product teams experiment.",
                  "body":"I studied your onboarding problem and want to build better workflows.",
                  "approved":True}).status_code==200
        preview=client.get(path+"/documents/preview?kind=cv")
        assert preview.status_code==200 and "Jordan Example" in preview.text
        for kind in ("cv","letter"):
            for fmt in ("html","pdf","docx"):
                r=client.post(path+"/documents/export",headers={"Origin":BASE},
                              json={"kind":kind,"format":fmt,"approved":True})
                assert r.status_code==200,(kind,fmt,r.text[:300])
                assert r.headers["cache-control"]=="private, no-store"
                if fmt=="pdf": assert len(PdfReader(io.BytesIO(r.content)).pages)==1
                if fmt=="docx":assert len(Document(io.BytesIO(r.content)).paragraphs)>0
        client.cookies.set("rh_session",other_token)
        assert client.get(path+"/documents/preview?kind=cv").status_code==404
        assert client.get("/web-api/editorial-profile").json()["exists"] is False
        assert client.post(path+"/documents/export",headers={"Origin":BASE},
            json={"kind":"cv","format":"pdf","approved":True}).status_code==404
